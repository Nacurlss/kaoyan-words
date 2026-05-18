#!/usr/bin/env python3
"""从原始试卷提取完形填空选项词 | Extract cloze test option words from raw exam files.

扫描全文（段落 + DOCX 表格），匹配 A/B/C/D 选项行，提取备选词，
写入 data/processed/sections/*/01_完形填空_options.txt。

Usage:
    python scripts/extract_cloze_options.py
"""


import sys
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src" / "backend"))

from parser import parse_docx, parse_pdf, parse_file


RAW_DIR = ROOT / "data" / "raw" / "exam_papers"
SECTIONS_DIR = ROOT / "data" / "processed" / "sections"

OPTION_LINE_RE = re.compile(
    r'\d{1,2}\s*[\.\s]\s*'
    r'\[A\][^\[]*?'
    r'\[B\][^\[]*?'
    r'\[C\][^\[]*?'
    r'\[D\][^\[]*',
    re.IGNORECASE,
)

OPTION_WORD_RE = re.compile(
    r'[\[［（(][A-Da-d][\]］）)]\s*([\w\-]+(?:\s+[\w\-]+)?)'
)


def extract_option_words(text: str) -> list[str]:
    """Extract unique option words from raw text lines like:
    '1. [A] affected [B] achieved [C] extracted [D] restored'
    """
    words = []
    seen = set()
    for line in text.split('\n'):
        stripped = line.strip()
        if not OPTION_LINE_RE.match(stripped):
            continue
        for m in OPTION_WORD_RE.finditer(line):
            word = m.group(1).lower().strip()
            if word and word not in seen:
                seen.add(word)
                words.append(word)
    return words


def read_docx_with_tables(filepath: Path) -> str:
    """Read DOCX including both paragraphs and table cells."""
    import docx as dx
    try:
        doc = dx.Document(str(filepath))
    except Exception:
        return parse_file(str(filepath))
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(' '.join(cells))
    return '\n'.join(lines)


def _extract_year(filename: str) -> str:
    m = re.search(r'(\d{4})', filename)
    return m.group(1) if m else "unknown"


def _detect_exam_type(text: str) -> str:
    if '英语二' in text:
        return '英语二'
    return '英语一'


def main():
    if not RAW_DIR.exists():
        print(f"Raw exam papers dir not found: {RAW_DIR}")
        return

    files = sorted(RAW_DIR.glob("*"))
    files = [f for f in files if f.suffix.lower() in (".pdf", ".docx", ".doc")]

    total_written = 0
    total_words = 0

    for filepath in files:
        suffix = filepath.suffix.lower()
        year = _extract_year(filepath.name)
        print(f"Processing: {filepath.name} (year={year})", end=" ")

        try:
            if suffix == ".pdf":
                text = parse_pdf(str(filepath))
            else:
                text = read_docx_with_tables(filepath)
        except Exception as e:
            print(f"→ ERROR: {e}")
            continue

        exam_type = _detect_exam_type(text)
        words = extract_option_words(text)
        if not words:
            print(f"→ no options found")
            continue

        target_dir = SECTIONS_DIR / f"{year}_{exam_type}"
        if not target_dir.exists():
            print(f"→ SKIP (dir not found: {target_dir})")
            continue

        out_path = target_dir / "01_完形填空_options.txt"
        out_path.write_text("\n".join(words) + "\n", encoding="utf-8")
        total_words += len(words)
        total_written += 1
        print(f"→ {len(words)} words → {out_path.name}")

    print(f"\nDone. Written {total_written} files, {total_words} total option words.")


if __name__ == "__main__":
    main()
